"""outline.json -> pptx（16:9，标题页 + 要点页）。用法：build.py <outline.json> <out.pptx>"""
import json
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

FONT = "微软雅黑"


def set_font(text_frame, size):
    for para in text_frame.paragraphs:
        for run in para.runs:
            run.font.name = FONT
            run.font.size = Pt(size)


def main(outline_path, out_path):
    with open(outline_path, encoding="utf-8") as f:
        outline = json.load(f)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = outline["title"]
    set_font(cover.shapes.title.text_frame, 40)
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = outline.get("subtitle", "")
        set_font(cover.placeholders[1].text_frame, 20)

    for spec in outline["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = spec["title"]
        set_font(slide.shapes.title.text_frame, 32)
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(spec.get("bullets", [])):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = bullet
        set_font(body, 20)

    prs.save(out_path)
    print(f"OK {out_path} slides={len(prs.slides)}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
