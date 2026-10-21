"""pptx -> PDF -> PNG，并做文字溢出的粗略检查（替代多模态视觉自检）。用法：preview.py <file.pptx>"""
import pathlib
import subprocess
import sys

from pptx import Presentation
from pptx.util import Pt

# 粗略估算：中文字符宽≈字号，英文≈0.55×字号
def text_width_pt(text, size_pt):
    return sum(size_pt if ord(c) > 0x2E80 else size_pt * 0.55 for c in text)


def check_overflow(pptx_path):
    problems = []
    for idx, slide in enumerate(Presentation(pptx_path).slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame or shape.width is None:
                continue
            box_w = shape.width / 12700  # EMU -> pt
            box_h = shape.height / 12700
            used_h = 0.0
            for para in shape.text_frame.paragraphs:
                size = next((r.font.size.pt for r in para.runs if r.font.size), 18)
                lines = max(1, int(text_width_pt(para.text, size) // max(box_w - 14, 1)) + 1)
                used_h += lines * size * 1.2
            if used_h > box_h:
                problems.append(f"OVERFLOW slide={idx} shape={shape.name} need={used_h:.0f}pt box={box_h:.0f}pt")
    return problems


def render(pptx_path):
    src = pathlib.Path(pptx_path)
    outdir = pathlib.Path("/workspace/preview")
    outdir.mkdir(exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(src)],
        check=True, timeout=180, capture_output=True,
    )
    pdf = outdir / (src.stem + ".pdf")
    subprocess.run(["pdftoppm", "-png", "-r", "50", str(pdf), str(outdir / src.stem)], check=True, timeout=120)
    return sorted(outdir.glob(src.stem + "*.png"))


if __name__ == "__main__":
    path = sys.argv[1]
    for png in render(path):
        print(f"PNG {png}")
    issues = check_overflow(path)
    print("\n".join(issues) if issues else "LAYOUT_OK")
