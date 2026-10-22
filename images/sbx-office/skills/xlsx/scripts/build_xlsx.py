"""Build a formatted .xlsx from a JSON workbook spec, or use the helpers from your own openpyxl code.

    python build_xlsx.py book.json out.xlsx      build
    python build_xlsx.py --example               print an example spec

    # in your own script (pandas analysis etc.):
    import sys; sys.path.insert(0, "/opt/skills/xlsx/scripts")
    from build_xlsx import write_table, add_chart, title_block, autofit

Cells whose value is a string starting with "=" are formulas, written as formulas (not computed values),
so the workbook stays live when people change the inputs. In formulas, {r} is the cell's own row and
{first} / {last} the first / last data row of the table, e.g. "=SUM(B{r}:D{r})" or "=E{r}/SUM(E{first}:E{last})". The platform recalculates with LibreOffice
for the preview; errors like #DIV/0! or #REF! show up there - fix them.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "微软雅黑"
THEMES = {  # header fill, header text, zebra fill, accent (titles, charts), chart palette
    "azure": ("1F5AA6", "FFFFFF", "EEF3FB", "1F5AA6", ["1F5AA6", "F2A93B", "4FB3BF", "8C6BB1", "E4572E"]),
    "emerald": ("0F766E", "FFFFFF", "E9F5F2", "0F766E", ["0F766E", "F59E0B", "3B82F6", "A855F7", "EF4444"]),
    "graphite": ("1F2937", "FFFFFF", "F1F3F6", "2563EB", ["2563EB", "22D3EE", "F59E0B", "10B981", "F43F5E"]),
    "crimson": ("9F1D17", "FFFFFF", "FBEFEE", "9F1D17", ["9F1D17", "C9A227", "3F6E8C", "6B8E23", "8E5572"]),
    "violet": ("5B21B6", "FFFFFF", "F3EEFD", "5B21B6", ["5B21B6", "EC4899", "06B6D4", "F59E0B", "10B981"]),
}
FORMATS = {  # friendly names for number formats
    "int": "#,##0", "number": "#,##0.00", "money": "¥#,##0.00", "yuan": "¥#,##0", "usd": "$#,##0.00",
    "percent": "0.0%", "pct": "0%", "date": "yyyy-mm-dd", "month": "yyyy-mm", "text": "@",
}
THIN = Side(style="thin", color="D9DEE6")


def _fmt(name: str | None) -> str | None:
    return FORMATS.get(name, name) if name else None


def _width(value) -> float:
    s = str(value if value is not None else "")
    if s.startswith("="):
        return 10
    return sum(2.0 if ord(c) > 0x2E80 else 1.1 for c in s)


def title_block(ws, title: str, subtitle: str = "", theme: str = "azure", width_cols: int = 6) -> int:
    """Title in A1 (and a subtitle in A2). Returns the first free row."""
    accent = THEMES.get(theme, THEMES["azure"])[3]
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=accent)
    ws.row_dimensions[1].height = 28
    row = 2
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(name=FONT, size=10, color="667085")
        row = 3
    return row + 1


def write_table(ws, top: int, left: int, columns: list, rows: list[list], theme: str = "azure", total: dict | None = None,
                zebra: bool = True, databar: str | None = None) -> tuple[int, int, int, int]:
    """Header + rows with formats. `columns` items: "name" or {"name", "format", "width", "align"}.
    `total`: {"label": "合计", "sum": ["销售额", ...]} adds a SUM row. Returns (top, left, bottom, right)."""
    head_fill, head_text, zebra_fill, _, _ = THEMES.get(theme, THEMES["azure"])
    cols = [c if isinstance(c, dict) else {"name": c} for c in columns]
    for j, c in enumerate(cols):
        cell = ws.cell(row=top, column=left + j, value=c["name"])
        cell.font = Font(name=FONT, bold=True, color=head_text)
        cell.fill = PatternFill("solid", fgColor=head_fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    ws.row_dimensions[top].height = 24
    first, last = top + 1, top + len(rows)
    for i, r in enumerate(rows, start=1):
        for j, c in enumerate(cols):
            v = r[j] if j < len(r) else None
            if isinstance(v, str) and v.startswith("="):
                v = v.replace("{r}", str(top + i)).replace("{first}", str(first)).replace("{last}", str(last))
            cell = ws.cell(row=top + i, column=left + j, value=v)
            cell.font = Font(name=FONT, size=10.5)
            cell.border = Border(bottom=THIN)
            fmt = _fmt(c.get("format"))
            if fmt:
                cell.number_format = fmt
            numeric = isinstance(v, int | float) or (isinstance(v, str) and v.startswith("="))
            cell.alignment = Alignment(horizontal=c.get("align") or ("right" if numeric else "left"), vertical="center")
            if zebra and i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=zebra_fill)
    bottom = top + len(rows)
    if total:
        bottom += 1
        ws.cell(row=bottom, column=left, value=total.get("label", "合计"))
        for j, c in enumerate(cols):
            cell = ws.cell(row=bottom, column=left + j)
            if c["name"] in total.get("sum", []):
                letter = get_column_letter(left + j)
                cell.value = f"=SUM({letter}{top + 1}:{letter}{bottom - 1})"
                cell.number_format = _fmt(c.get("format")) or "#,##0"
                cell.alignment = Alignment(horizontal="right")
            cell.font = Font(name=FONT, bold=True)
            cell.border = Border(top=Side(style="medium", color=head_fill))
    right = left + len(cols) - 1
    for j, c in enumerate(cols):
        letter = get_column_letter(left + j)
        widest = max([_width(c["name"])] + [_width(r[j] if j < len(r) else "") for r in rows[:200]])
        ws.column_dimensions[letter].width = c.get("width") or min(48, max(8, widest + 3))
    if databar:
        j = next((k for k, c in enumerate(cols) if c["name"] == databar), None)
        if j is not None:
            letter = get_column_letter(left + j)
            ws.conditional_formatting.add(f"{letter}{top + 1}:{letter}{top + len(rows)}",
                                          DataBarRule(start_type="min", end_type="max", color=THEMES.get(theme, THEMES["azure"])[3]))
    ws.freeze_panes = ws.cell(row=top + 1, column=left)
    return top, left, bottom, right


def add_chart(ws, kind: str, table: tuple[int, int, int, int], value_cols: list[int], anchor: str, title: str = "",
              theme: str = "azure", category_col: int = 0, data_rows: int | None = None, y_title: str = "") -> None:
    """Native Excel chart from a table written by write_table. value_cols / category_col are 0-based column
    offsets inside the table; data_rows limits how many rows are charted (e.g. to leave out the total row)."""
    top, left, bottom, _ = table
    last = top + data_rows if data_rows else bottom
    palette = THEMES.get(theme, THEMES["azure"])[4]
    chart = {"bar": BarChart, "column": BarChart, "line": LineChart, "pie": PieChart}.get(kind, BarChart)()
    if kind == "bar":
        chart.type = "bar"
    elif kind == "column":
        chart.type = "col"
    cats = Reference(ws, min_col=left + category_col, min_row=top + 1, max_row=last)
    for k in value_cols:
        data = Reference(ws, min_col=left + k, min_row=top, max_row=last)
        chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.title = title or None
    chart.height, chart.width = 8, 16
    if kind != "pie":
        for i, s in enumerate(chart.series):
            color = palette[i % len(palette)]
            if kind == "line":
                s.graphicalProperties.line.solidFill = color
                s.graphicalProperties.line.width = 28000
                s.smooth = False
            else:
                s.graphicalProperties.solidFill = color
                s.graphicalProperties.line.solidFill = color
        if y_title:
            chart.y_axis.title = y_title
        chart.legend.position = "b"
        if len(value_cols) == 1:
            chart.legend = None
    ws.add_chart(chart, anchor)


def autofit(ws) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        widest = max((_width(c.value) for c in col[:300]), default=8)
        ws.column_dimensions[letter].width = min(48, max(8, widest + 2))


def build(spec: dict, out: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    theme = spec.get("theme", "azure")
    for sh in spec.get("sheets", []):
        ws = wb.create_sheet(str(sh.get("name", "Sheet"))[:31])
        ws.sheet_view.showGridLines = False
        row = title_block(ws, sh["title"], sh.get("subtitle", ""), theme) if sh.get("title") else 1
        table = write_table(ws, row, 1, sh.get("columns", []), sh.get("rows", []), theme, total=sh.get("total"), databar=sh.get("databar"))
        n_rows = len(sh.get("rows", []))
        for k, ch in enumerate(sh.get("charts", [])):
            names = [c if isinstance(c, str) else c["name"] for c in sh.get("columns", [])]
            cols = [names.index(v) for v in ch.get("values", []) if v in names]
            cat = names.index(ch["category"]) if ch.get("category") in names else 0
            anchor = ch.get("anchor") or f"{get_column_letter(table[3] + 2)}{row + k * 17}"
            add_chart(ws, ch.get("type", "column"), table, cols, anchor, ch.get("title", ""), theme, cat, data_rows=n_rows,
                      y_title=ch.get("y_title", ""))
        for note in sh.get("notes", []):
            r = ws.max_row + 2
            ws.cell(row=r, column=1, value=note).font = Font(name=FONT, size=9, color="667085")
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


EXAMPLE = {
    "theme": "azure",
    "sheets": [
        {
            "name": "销售汇总", "title": "2026 年三季度销售汇总", "subtitle": "单位：元；数据截至 9 月 30 日",
            "columns": ["区域", {"name": "7月", "format": "yuan"}, {"name": "8月", "format": "yuan"}, {"name": "9月", "format": "yuan"},
                        {"name": "季度合计", "format": "yuan"}, {"name": "环比", "format": "percent"}],
            "rows": [
                ["华东", 1200000, 1350000, 1500000, "=SUM(B{r}:D{r})", "=D{r}/C{r}-1"],
                ["华南", 980000, 1010000, 990000, "=SUM(B{r}:D{r})", "=D{r}/C{r}-1"],
                ["华北", 760000, 820000, 905000, "=SUM(B{r}:D{r})", "=D{r}/C{r}-1"],
            ],
            "total": {"label": "合计", "sum": ["7月", "8月", "9月", "季度合计"]},
            "databar": "季度合计",
            "charts": [{"type": "column", "title": "各区域季度销售额", "category": "区域", "values": ["季度合计"]}],
            "notes": ["说明：环比 = 9 月 / 8 月 - 1"],
        }
    ],
}


def main(argv: list[str]) -> int:
    if len(argv) == 2 and argv[1] == "--example":
        print(json.dumps(EXAMPLE, ensure_ascii=False, indent=2))
        return 0
    if len(argv) != 3:
        print(__doc__)
        return 2
    build(json.loads(Path(argv[1]).read_text(encoding="utf-8")), argv[2])
    print(f"OK {argv[2]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
