"""Diagnose the pptx geometry and overflow heuristic (dev only)."""

import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "images/sbx-office/skills/pptx-builder/scripts"))
from build import set_font
from preview import check_overflow


def build(path):
    from pptx.util import Inches

    outline = {
        "title": "AILab 简易版冒烟测试",
        "subtitle": "gVisor 沙箱",
        "slides": [
            {
                "title": "第一页：中文排版",
                "bullets": ["微软雅黑映射到 Noto Sans CJK", "LibreOffice 渲染预览"],
            },
            {
                "title": "第二页：故意溢出",
                "bullets": [
                    "这是一条非常非常长的要点，用来测试溢出检查是否能够发现文字超出文本框的情况，"
                    "重复一遍：这是一条非常非常长的要点，用来测试溢出检查是否能够发现文字超出文本框的情况",
                    "a",
                    "b",
                    "c",
                    "d",
                    "e",
                    "f",
                    "g",
                    "h",
                ],
            },
        ],
    }
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = outline["title"]
    set_font(cover.shapes.title.text_frame, 40)
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = outline.get("subtitle", "")
        set_font(cover.placeholders[1].text_frame, 20)
    for spec in outline["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = spec["title"]
        set_font(slide.shapes.title.text_frame, 32)
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(spec.get("bullets", [])):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = bullet
        set_font(body, 20)
    prs.save(path)


def main(path):
    build(path)
    prs = Presentation(path)
    print("slide size pt:", prs.slide_width / 12700, prs.slide_height / 12700)
    for i, slide in enumerate(prs.slides, 1):
        print(f"--- slide {i}")
        for sh in slide.shapes:
            print(
                f"  {sh.shape_type} {sh.name!r} L={sh.left / 12700:.1f} T={sh.top / 12700:.1f} "
                f"W={sh.width / 12700:.1f} H={sh.height / 12700:.1f}"
            )
            if sh.has_text_frame:
                print("    paras:", [p.text[:20] for p in sh.text_frame.paragraphs])
    print("overflow:", check_overflow(path))


if __name__ == "__main__":
    main("/tmp/diag.pptx")
