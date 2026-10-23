"""File processing pipeline: scan, extract text, write derivatives."""

from __future__ import annotations

import io
import json
import logging
import warnings
from pathlib import PurePath

from PIL import Image, ImageOps

from .. import skillfiles
from ..clients import PlatformClient
from ..config import settings
from ..storage import Storage

log = logging.getLogger("ailab.worker")

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".xml", ".html", ".log", ".py", ".sh"}
MAX_TEXT = 200_000
# Bound both compressed input and decoded pixels; parsing runs only in the Worker.
MAX_THUMBNAIL_BYTES = 20 * 1024 * 1024
# Formats decoded at full size (PNG, GIF, WebP...) must fit the Worker's memory. JPEG decodes at down to 1/8
# scale (draft), so camera photos of 24-48 MP+ cost ~2 MP of pixels; a 16 MP cap for them would skip most photos.
MAX_DECODED_PIXELS = 16_000_000
MAX_JPEG_PIXELS = 128_000_000
Image.MAX_IMAGE_PIXELS = MAX_JPEG_PIXELS


def _thumbnail(file_id: str, data: bytes, meta: dict, storage: Storage) -> dict:
    if (
        len(data) > MAX_THUMBNAIL_BYTES
        or (meta.get("content_type") or "").split(";", 1)[0].lower() == "image/svg+xml"
        or PurePath(meta["filename"]).suffix.lower() == ".svg"
    ):
        return {}
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            # Pillow detects raster formats even if the uploader supplied a generic MIME type.
            with Image.open(io.BytesIO(data)) as source:
                if source.format in {"EPS", "WMF", "PDF"}:
                    return {}  # don't launch vector rasterizers such as Ghostscript
                width, height = source.size
                if width * height > (MAX_JPEG_PIXELS if source.format == "JPEG" else MAX_DECODED_PIXELS):
                    return {}
                source.seek(0)  # animations use their first frame
                # JPEG can downsample during decoding, before allocating the full raster.
                source.draft("RGB", (480, 480))
                # A square bound commutes with EXIF rotations/reflections. Shrink first so
                # exif_transpose never copies a full decoded 16 MP frame.
                source.thumbnail((480, 480))
                with ImageOps.exif_transpose(source) as oriented:
                    mode = "RGBA" if "A" in oriented.getbands() or "transparency" in oriented.info else "RGB"
                    with oriented.convert(mode) as converted:
                        output = io.BytesIO()
                        converted.save(output, "WEBP", quality=75, method=4)
        key = f"derived/{file_id}/thumb.webp"
        storage.put_bytes(key, output.getvalue(), "image/webp")
        return {"thumb_key": key}
    except Exception:
        # Decoding/encoding/storage failures must never change the file's usability.
        log.debug("thumbnail unavailable for %s", file_id, exc_info=True)
        return {}


def scan_bytes(data: bytes) -> dict:
    """Placeholder for ClamAV/ICAP; catches the EICAR test string."""
    if b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE" in data:
        return {"ok": False, "engine": "stub", "reason": "EICAR test signature"}
    return {"ok": True, "engine": "stub"}


def _pdf(data: bytes) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = len(reader.pages)
    text = "\n".join((page.extract_text() or "") for page in reader.pages[:20])
    return {"kind": "pdf", "pages": pages, "chars": len(text), "text": text[:MAX_TEXT]}


def _docx(data: bytes) -> dict:
    import docx

    doc = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs if p.text)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + "\t".join(cell.text for cell in row.cells)
    return {"kind": "docx", "paragraphs": len(doc.paragraphs), "chars": len(text), "text": text[:MAX_TEXT]}


def _pptx(data: bytes) -> dict:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    chunks: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                chunks.append(shape.text)
    text = "\n".join(chunks)
    return {"kind": "pptx", "slides": len(prs.slides), "chars": len(text), "text": text[:MAX_TEXT]}


def _xlsx(data: bytes) -> dict:
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - optional dependency
        return {"kind": "xlsx", "sheets": 0, "chars": 0, "text": ""}
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks: list[str] = []
    for ws in wb.worksheets:
        chunks.append(f"# {ws.title}")
        for row in ws.iter_rows(max_row=200, values_only=True):
            chunks.append("\t".join("" if v is None else str(v) for v in row))
    text = "\n".join(chunks)
    return {"kind": "xlsx", "sheets": len(wb.worksheets), "chars": len(text), "text": text[:MAX_TEXT]}


def extract(data: bytes, filename: str, content_type: str) -> dict:
    suffix = PurePath(filename).suffix.lower()
    try:
        if suffix == ".pdf" or content_type == "application/pdf":
            return _pdf(data)
        if suffix == ".docx":
            return _docx(data)
        if suffix == ".pptx":
            return _pptx(data)
        if suffix == ".xlsx":
            return _xlsx(data)
    except Exception as exc:
        log.warning("extract failed for %s: %s", filename, exc)
        return {"kind": "error", "error": f"{type(exc).__name__}: {exc}", "chars": 0, "text": ""}

    if suffix in TEXT_SUFFIXES or content_type.startswith("text/"):
        text = data.decode("utf-8", "replace")
        return {"kind": "text", "chars": len(text), "text": text[:MAX_TEXT]}

    return {"kind": "binary", "chars": 0, "text": "", "bytes": len(data)}


def process_file(
    file_id: str, platform: PlatformClient | None = None, storage: Storage | None = None, *, thumbnail_only: bool = False
) -> dict:
    if platform is not None:
        return _process_file(file_id, platform, storage, thumbnail_only=thumbnail_only)
    with PlatformClient() as owned:  # holds a connection pool: close it with the file
        return _process_file(file_id, owned, storage, thumbnail_only=thumbnail_only)


def _process_file(file_id: str, platform: PlatformClient, storage: Storage | None, *, thumbnail_only: bool) -> dict:
    from ..storage import get_storage

    storage = storage or get_storage()
    meta = platform.get_file(file_id)
    if thumbnail_only:
        if meta.get("status") != "ready" or (meta.get("meta") or {}).get("thumb_key"):
            return {"file_id": file_id, "skipped": meta.get("status")}
        if meta.get("size", 0) > MAX_THUMBNAIL_BYTES:
            return {"file_id": file_id, "skipped": "too_large"}
        data = platform.get_file_bytes(file_id, max_bytes=MAX_THUMBNAIL_BYTES)
        thumb = _thumbnail(file_id, data, meta, storage)
        if thumb:
            platform.patch_file(file_id, meta=thumb)
        return {"file_id": file_id, **thumb}
    if meta.get("status") != "pending":
        return {"file_id": file_id, "skipped": meta.get("status")}
    data = platform.get_file_bytes(file_id)

    scan = scan_bytes(data)
    if not scan["ok"]:
        platform.patch_file(file_id, meta={"processing": "rejected", "scan": scan}, status="rejected")
        return {"file_id": file_id, "rejected": scan}

    thumb = _thumbnail(file_id, data, meta, storage)

    if len(data) > settings.max_parse_bytes:
        # Office/PDF parsers can use several times the file size. Keep the file
        # usable but skip the parse instead of OOMing the 256 MiB Worker pod.
        info = {"kind": "too_large", "chars": 0, "bytes": len(data), "truncated": True}
        text = ""
    else:
        info = extract(data, meta["filename"], meta.get("content_type") or "")
        text = info.pop("text", "")
    skill: dict = {}
    try:  # a personal skill (SKILL.md, or a zip with one): recognised here, listed on the 技能 page
        found = skillfiles.inspect(data, meta["filename"])
        if found:
            skill = {"skill": found}
    except skillfiles.SkillFileError as exc:
        skill = {"skill_error": str(exc)}
    derived = {
        "file_id": file_id,
        "filename": meta["filename"],
        "content_type": meta.get("content_type"),
        **info,
        "excerpt": text[:1000],
    }
    key = f"derived/{file_id}/meta.json"
    if text:  # the whole extracted text, so a conversation can read an attachment without a sandbox
        derived["text_key"] = f"derived/{file_id}/text.txt"
        storage.put_bytes(derived["text_key"], text.encode("utf-8"), "text/plain; charset=utf-8")
    storage.put_bytes(key, json.dumps(derived, ensure_ascii=False, indent=2).encode(), "application/json")

    platform.patch_file(
        file_id,
        meta={
            "processing": "done",
            "scan": scan,
            "extract": info,
            "derived_key": key,
            **({"text_key": derived["text_key"]} if text else {}),
            "excerpt": text[:300],
            **skill,
            **thumb,
        },
        status="ready",
    )
    return {"file_id": file_id, "kind": info.get("kind"), "chars": info.get("chars", 0), "derived_key": key}
