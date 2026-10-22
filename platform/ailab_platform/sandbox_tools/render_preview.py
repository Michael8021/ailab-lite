"""Render a preview of an office file. Runs INSIDE the sandbox (copied there by the Broker).

    python render_preview.py <file> <out_dir>

Prints one JSON line: {"kind": "pages", "pages": [names...]} for slides / documents / PDFs (JPEG per page)
or {"kind": "sheets", "data": "sheets.json"} for spreadsheets, or {"error": "..."}.
Only standard tools of the sandbox image are used: LibreOffice, pdftoppm, openpyxl.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

PAGES = {".pptx", ".ppt", ".odp", ".docx", ".doc", ".odt", ".pdf"}
SHEETS = {".xlsx", ".xlsm", ".xls", ".ods", ".csv"}
MAX_PAGES = 40
MAX_ROWS, MAX_COLS = 200, 40


def soffice(src: Path, fmt: str, outdir: Path) -> Path:
    profile = Path(tempfile.mkdtemp(prefix="lo-profile-"))  # a private profile: parallel runs do not collide
    try:
        subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", fmt, "--outdir", str(outdir), str(src)],
            check=True,
            capture_output=True,
            timeout=180,
        )
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    out = outdir / f"{src.stem}.{fmt.split(':')[0]}"
    if not out.exists():
        raise RuntimeError(f"LibreOffice produced no {fmt} for {src.name}")
    return out


def pages(src: Path, out: Path) -> dict:
    with tempfile.TemporaryDirectory() as tmp:
        pdf = src if src.suffix.lower() == ".pdf" else soffice(src, "pdf", Path(tmp))
        subprocess.run(
            ["pdftoppm", "-jpeg", "-jpegopt", "quality=82", "-scale-to", "1280", "-l", str(MAX_PAGES), str(pdf), str(out / "p")],
            check=True,
            capture_output=True,
            timeout=180,
        )
    names = sorted((p.name for p in out.glob("p-*.jpg")), key=lambda n: int(n[2:-4]))
    return {"kind": "pages", "pages": names}


def _color(c) -> str | None:
    try:
        if c is not None and c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) == 8 and c.rgb[:2] != "00":
            return c.rgb[2:]
    except Exception:
        pass
    return None


def _display(v, fmt: str) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int | float):
        f = fmt or ""
        if "%" in f:
            decimals = len(f.split(".")[1].split("%")[0]) if "." in f else 0
            return f"{v * 100:.{decimals}f}%"
        prefix = "¥" if "¥" in f else "$" if "$" in f else ""
        if "#,##0" in f or prefix:
            decimals = len(f.split(".")[1].rstrip(");-_ ")) if "." in f.split(";")[0] else 0
            decimals = min(decimals, 4)
            return f"{prefix}{v:,.{decimals}f}"
        if isinstance(v, float):
            return f"{v:.6g}"
        return str(v)
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d") if "h" not in (fmt or "") else v.strftime("%Y-%m-%d %H:%M")
    return str(v)


def sheets(src: Path, out: Path) -> dict:
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory() as tmp:
        native = src.suffix.lower() in (".xlsx", ".xlsm")
        if shutil.which("soffice"):
            # LibreOffice recalculates on load and stores the results, so formula cells get values (openpyxl alone cannot)
            recalced = soffice(src, "xlsx:Calc MS Excel 2007 XML", Path(tmp))
        elif native:
            recalced = src  # no LibreOffice (a dev machine): formulas are shown as written, without values
        else:
            raise RuntimeError("LibreOffice is needed to preview this kind of spreadsheet")
        values = load_workbook(recalced, data_only=True)
        styles = load_workbook(src if native else recalced)
        result = []
        errors = []
        for ws in values.worksheets:
            st = styles[ws.title] if ws.title in styles.sheetnames else ws
            nrows, ncols = min(ws.max_row, MAX_ROWS), min(ws.max_column, MAX_COLS)
            rows = []
            for r in range(1, nrows + 1):
                row = []
                for c in range(1, ncols + 1):
                    cell, sc = ws.cell(r, c), st.cell(r, c)
                    text = _display(cell.value, sc.number_format)
                    if not text and isinstance(sc.value, str) and sc.value.startswith("=") and recalced is src:
                        text = sc.value
                    if text.startswith("#") and text.rstrip("!?0/") in ("#DIV", "#REF", "#VALUE", "#NAME", "#N/A", "#NUM", "#NULL"):
                        errors.append(f"{ws.title}!{cell.coordinate} {text}")
                    item = {"v": text}
                    if sc.font is not None and sc.font.bold:
                        item["b"] = 1
                    if sc.fill is not None and sc.fill.fill_type == "solid" and (bg := _color(sc.fill.fgColor)):
                        item["bg"] = bg
                    if sc.font is not None and (fg := _color(sc.font.color)):
                        item["fg"] = fg
                    if isinstance(cell.value, int | float) and not isinstance(cell.value, bool):
                        item["n"] = 1
                    if isinstance(sc.value, str) and sc.value.startswith("="):
                        item["f"] = sc.value[:200]
                    row.append(item)
                rows.append(row)
            widths = []
            for c in range(1, ncols + 1):
                from openpyxl.utils import get_column_letter

                dim = st.column_dimensions.get(get_column_letter(c))
                widths.append(round(dim.width, 1) if dim is not None and dim.width else None)
            result.append(
                {
                    "name": ws.title,
                    "rows": rows,
                    "widths": widths,
                    "merged": [str(m) for m in st.merged_cells.ranges][:200],
                    "charts": len(getattr(st, "_charts", [])),
                    "truncated": ws.max_row > MAX_ROWS or ws.max_column > MAX_COLS,
                    "size": [ws.max_row, ws.max_column],
                }
            )
    (out / "sheets.json").write_text(json.dumps({"sheets": result, "errors": errors[:50]}, ensure_ascii=False), encoding="utf-8")
    return {"kind": "sheets", "data": "sheets.json", "errors": errors[:20]}


def main() -> int:
    src, out = Path(sys.argv[1]), Path(sys.argv[2])
    out.mkdir(parents=True, exist_ok=True)
    ext = src.suffix.lower()
    try:
        if ext in PAGES:
            result = pages(src, out)
        elif ext in SHEETS:
            result = sheets(src, out)
        else:
            result = {"error": f"no preview for {ext} files"}
    except subprocess.TimeoutExpired:
        result = {"error": "rendering timed out"}
    except Exception as exc:
        result = {"error": f"{type(exc).__name__}: {exc}"[:500]}
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
